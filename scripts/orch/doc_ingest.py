#!/usr/bin/env python3
"""ADF document ingest — turn an uploaded requirement source (doc / data sample) into
extracted, structured requirement text for the crew's Wave-1 ingestion. Uploaded docs
are the user's STATED requirements, so they are weighted above web research and the PRD
must trace back to them.

Formats (dependency-light; a missing library yields a NOTE, never a crash):
  .md / .markdown / .txt  -> read text (stdlib)
  .csv                    -> profile columns + sample rows (stdlib `csv`) → data-model facts
  .xlsx                   -> `openpyxl` if present, else degrade
  .pdf                    -> `pdfplumber`/`PyPDF2` if present, else degrade (caller may
                             fall back to Claude native-PDF)
  .docx                   -> `python-docx` if present, else degrade
  (images/mockups are routed to a vision model by the crew, not here)

Then, optionally, a model (`complete(prompt, role)`) extracts structured requirements
from the raw text — grounding the draft in the user's actual document.

  extract_text(path)            -> {kind, text, note}        (text='' + note on degrade)
  ingest(path, complete)        -> {source, kind, requirements, raw}
"""
import csv
import io
import os


def _read(path):
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _csv_profile(path):
    """A compact data-model profile from a CSV: header + up to 5 sample rows. Grounds
    'data-entities' requirements in the user's real columns."""
    try:
        with open(path, newline="", encoding="utf-8", errors="replace") as f:
            rows = list(csv.reader(f))
    except OSError as e:
        return "", f"unreadable CSV: {e}"
    if not rows:
        return "", "empty CSV"
    header, sample = rows[0], rows[1:6]
    buf = io.StringIO()
    buf.write("Columns: " + ", ".join(header) + "\n")
    buf.write(f"Rows: {len(rows) - 1}\nSample:\n")
    for r in sample:
        buf.write("  " + " | ".join(r) + "\n")
    return buf.getvalue(), ""


def _xlsx_profile(path):
    try:
        import openpyxl  # noqa: F401
    except Exception:
        return "", "xlsx ingest needs `openpyxl` (pip install openpyxl) — skipped"
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets[:3]:
            out.append(f"Sheet {ws.title}:")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > 5:
                    break
                out.append("  " + " | ".join("" if c is None else str(c) for c in row))
        return "\n".join(out), ""
    except Exception as e:  # noqa: BLE001
        return "", f"xlsx read failed: {e}"


def _pdf_text(path):
    for mod, fn in (("pdfplumber", "_pdfplumber"), ("PyPDF2", "_pypdf2")):
        try:
            __import__(mod)
        except Exception:
            continue
        return globals()[fn](path)
    return "", ("pdf ingest needs `pdfplumber` or `PyPDF2` — skipped (the crew may fall "
                "back to Claude native-PDF)")


def _pdfplumber(path):
    import pdfplumber
    try:
        with pdfplumber.open(path) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages), ""
    except Exception as e:  # noqa: BLE001
        return "", f"pdf read failed: {e}"


def _pypdf2(path):
    import PyPDF2
    try:
        r = PyPDF2.PdfReader(path)
        return "\n".join((pg.extract_text() or "") for pg in r.pages), ""
    except Exception as e:  # noqa: BLE001
        return "", f"pdf read failed: {e}"


def _docx_text(path):
    try:
        import docx
    except Exception:
        return "", "docx ingest needs `python-docx` (pip install python-docx) — skipped"
    try:
        d = docx.Document(path)
        return "\n".join(p.text for p in d.paragraphs), ""
    except Exception as e:  # noqa: BLE001
        return "", f"docx read failed: {e}"


def extract_text(path):
    """Raw text (+ a `kind` and a `note` on any degrade) from a document path."""
    ext = os.path.splitext(path)[1].lower()
    if not os.path.isfile(path):
        return {"kind": ext or "file", "text": "", "note": f"missing file: {path}"}
    if ext in (".md", ".markdown", ".txt", ""):
        try:
            return {"kind": "text", "text": _read(path), "note": ""}
        except OSError as e:
            return {"kind": "text", "text": "", "note": f"unreadable: {e}"}
    if ext == ".csv":
        text, note = _csv_profile(path)
        return {"kind": "data", "text": text, "note": note}
    if ext == ".xlsx":
        text, note = _xlsx_profile(path)
        return {"kind": "data", "text": text, "note": note}
    if ext == ".pdf":
        text, note = _pdf_text(path)
        return {"kind": "pdf", "text": text, "note": note}
    if ext == ".docx":
        text, note = _docx_text(path)
        return {"kind": "docx", "text": text, "note": note}
    return {"kind": ext, "text": "", "note": f"unsupported document type: {ext}"}


def ingest(path, complete=None):
    """Ingest a document into structured requirement text. With a model `complete`, it
    extracts the explicit requirements/constraints; otherwise returns the raw text.
    Always returns a dict (with a `note` on degrade) — never crashes the stage."""
    doc = extract_text(path)
    raw = (doc["text"] or "").strip()
    name = os.path.basename(path)
    if not raw:
        return {"source": name, "kind": doc["kind"], "requirements": "",
                "raw": "", "note": doc["note"]}
    requirements = raw[:8000]
    role = "draft" if doc["kind"] == "data" else "extract"
    if complete is not None:
        prompt = (
            "This is a user-provided requirements source — extract the EXPLICIT "
            "requirements, constraints, and acceptance criteria it states, as terse "
            f"bullet points (these are AUTHORITATIVE). Source: {name}\n\n{raw[:8000]}")
        res = complete(prompt, role)
        if res and res[0] and res[0].strip():
            requirements = res[0].strip()
    return {"source": name, "kind": doc["kind"], "requirements": requirements,
            "raw": raw[:8000], "note": doc["note"]}
